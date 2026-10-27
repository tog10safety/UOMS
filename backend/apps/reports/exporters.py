"""
Renders a (title, headers, rows) report tuple into PDF or XLSX bytes.
"""
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def build_pdf(title, headers, rows):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        topMargin=0.6 * inch, bottomMargin=0.6 * inch, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(title, styles['Title']),
        Paragraph(f'Generated {timezone.now():%Y-%m-%d %H:%M} UTC — UOMS', styles['Normal']),
        Spacer(1, 0.25 * inch),
    ]

    table_data = [headers] + [[str(cell) for cell in row] for row in rows]
    if len(table_data) == 1:
        table_data.append(['No records found.'] + [''] * (len(headers) - 1))

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#16213a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f7fa')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def build_xlsx(title, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]  # Excel sheet-name length limit

    header_fill = PatternFill(start_color='16213A', end_color='16213A', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, header in enumerate(headers, start=1):
        max_len = max([len(str(header))] + [len(str(r[col_idx - 1])) for r in rows] or [0])
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
