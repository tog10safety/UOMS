"""Excel import/export for the hazard registry (same column layout as the legacy system)."""
import io
import re

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import risk
from .models import Hazard

REQUIRED = [
    'Activity Code', 'Hazard Description', 'Before Likelihood', 'Before Severity',
    'Mitigation Strategies', 'After Likelihood', 'After Severity',
]
EXPORT_HEADERS = [
    'ID', 'Unit', 'Activity Code', 'Hazard Description', 'Before Likelihood', 'Before Severity',
    'Before Risk Rating', 'Mitigation Strategies', 'After Likelihood', 'After Severity',
    'After Risk Rating', 'Last Updated', 'Created By',
]
MAX_ROWS = 2000
_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean_text(value):
    """Strip control characters and Excel's `_x000D_` artefacts; normalise newlines."""
    if value is None:
        return ''
    text = str(value).replace('_x000D_', '')
    text = _ILLEGAL.sub('', text).replace('\r\n', '\n').replace('\r', '\n')
    return text.strip()


def _style_header(ws, headers, fill='1F4E78'):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _autosize(ws, cap=60):
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, cap)


def export_hazards(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hazard Registry'
    _style_header(ws, EXPORT_HEADERS)
    for row, h in enumerate(queryset.select_related('unit', 'created_by'), start=2):
        ws.append([
            str(h.id), h.unit.name, h.activity_code, h.hazard_description, h.before_likelihood, h.before_severity,
            h.before_risk_rating, h.mitigations, h.after_likelihood, h.after_severity, h.after_risk_rating,
            h.updated_at.strftime('%Y-%m-%d %H:%M'), h.created_by.full_name if h.created_by else '',
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hazards'
    _style_header(ws, REQUIRED)
    ws.append(['Ground Operations', 'Vehicle rollover on uneven terrain', 3, 'B',
               'Brief drivers\nInspect vehicle before departure', 2, 'C'])
    guide = wb.create_sheet('Instructions')
    _style_header(guide, ['Column', 'Description', 'Allowed values'], fill='BF8F00')
    for line in [
        ('Activity Code', 'Activity this hazard belongs to (any text).', 'Free text'),
        ('Hazard Description', 'What could go wrong.', 'Free text'),
        ('Before Likelihood', 'Likelihood before mitigation.', '1-5'),
        ('Before Severity', 'Severity before mitigation.', 'A-E'),
        ('Mitigation Strategies', 'One mitigation per line.', 'Free text'),
        ('After Likelihood', 'Likelihood after mitigation.', '1-5'),
        ('After Severity', 'Severity after mitigation.', 'A-E'),
    ]:
        guide.append(list(line))
    _autosize(ws)
    _autosize(guide, 80)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_hazards(file, unit, user, overwrite=False):
    """
    Returns {'created', 'updated', 'skipped', 'skip_reasons', 'errors'}.
    A hazard is a duplicate when unit + activity code + description already exist.
    """
    result = {'created': 0, 'updated': 0, 'skipped': 0, 'skip_reasons': [], 'errors': []}
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        result['errors'].append('The file is not a valid .xlsx workbook.')
        return result

    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = [str(c).strip() if c is not None else '' for c in next(rows, [])]
    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        result['errors'].append(f'Missing required columns: {", ".join(missing)}')
        return result

    def skip(row_number, reason):
        result['skipped'] += 1
        result['skip_reasons'].append(f'Row {row_number}: {reason}')

    with transaction.atomic():
        for row_number, row in enumerate(rows, start=2):
            if row_number - 1 > MAX_ROWS:
                result['errors'].append(f'Import limited to {MAX_ROWS} rows; the rest were ignored.')
                break
            if row is None or not any(c not in (None, '') for c in row):
                continue
            data = {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers)}
            activity = clean_text(data.get('Activity Code'))
            description = clean_text(data.get('Hazard Description'))
            if not activity or not description:
                skip(row_number, 'empty activity code or hazard description')
                continue
            mitigations = clean_text(data.get('Mitigation Strategies'))
            if not mitigations:
                skip(row_number, 'mitigation strategies are required')
                continue
            try:
                bl, bs = risk.normalize(data.get('Before Likelihood'), data.get('Before Severity'))
                al, a_s = risk.normalize(data.get('After Likelihood'), data.get('After Severity'))
            except risk.RiskInputError as exc:
                skip(row_number, str(exc))
                continue

            existing = Hazard.objects.filter(
                unit=unit, activity_code=activity, hazard_description=description, is_active=True,
            ).first()
            if existing:
                if not overwrite:
                    skip(row_number, f"hazard '{activity} - {description[:40]}' already exists (use overwrite to update)")
                    continue
                existing.before_likelihood, existing.before_severity = bl, bs
                existing.after_likelihood, existing.after_severity = al, a_s
                existing.mitigations = mitigations
                existing.save()
                result['updated'] += 1
            else:
                Hazard.objects.create(
                    unit=unit, created_by=user, activity_code=activity[:100], hazard_description=description,
                    before_likelihood=bl, before_severity=bs, mitigations=mitigations,
                    after_likelihood=al, after_severity=a_s,
                )
                result['created'] += 1
    return result
